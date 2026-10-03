"""Загрузка общего отчёта дистрибьютора за несколько месяцев с разбивкой по кварталам."""

from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .importer import _cell, _collect, parse_number, split_by_quarter
from .models import QUARTERS, Report

HEAD_FILL = PatternFill("solid", fgColor="1C1B19")
LABEL_FILL = PatternFill("solid", fgColor="FBE3DD")
ARTIST_FILL = PatternFill("solid", fgColor="DDF1EE")


def _fmt_rate(rate):
    rate = Decimal(rate)
    return f"{rate.normalize():f}".replace(".", ",")


def build_detail_xlsx(artist_name, year, quarter, header, rows, cols, rate):
    """Детализация за квартал: строки дистрибьютора + доля лейбла и доля артиста."""
    rate = Decimal(rate)
    share = rate / 100
    wb = Workbook()
    ws = wb.active
    ws.title = f"Q{quarter} {year}"
    head = [str(h) if h is not None else "" for h in header]
    label_col = f"Доля лейбла, {_fmt_rate(100 - rate)}%"
    artist_col = f"Вознаграждение артиста, {_fmt_rate(rate)}%"
    ws.append([f"{artist_name} — {dict(QUARTERS)[quarter]} {year}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Детализация начислений. «{head[cols['amount']]}» — доход лейбла по отчёту дистрибьютора; "
               f"артисту причитается {_fmt_rate(rate)}% по лицензионному договору."])
    ws["A2"].font = Font(italic=True, color="6B6860")
    ws.append([])
    ws.append(head + [label_col, artist_col])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    total_g = total_l = total_a = Decimal(0)
    total_q = 0
    for row in rows:
        values = list(row) + [None] * (len(head) - len(row))
        gross = parse_number(_cell(row, cols["amount"])) or Decimal(0)
        label = (gross * (1 - share)).quantize(Decimal("0.000001"))
        artist = (gross * share).quantize(Decimal("0.000001"))
        total_g += gross
        total_l += label
        total_a += artist
        total_q += int(parse_number(_cell(row, cols.get("quantity"))) or 0) if "quantity" in cols else 0
        ws.append(values[:len(head)] + [float(label), float(artist)])
    n = len(head)
    total_row = [None] * (n + 2)
    total_row[0] = "Итого"
    if "quantity" in cols:
        total_row[cols["quantity"]] = total_q
    total_row[cols["amount"]] = float(total_g)
    total_row[n] = float(total_l)
    total_row[n + 1] = float(total_a)
    ws.append(total_row)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    money_cols = [cols["amount"] + 1, n + 1, n + 2]
    for r in ws.iter_rows(min_row=hdr + 1, max_row=ws.max_row):
        for idx in money_cols:
            r[idx - 1].number_format = "#,##0.00"
        r[n].fill = LABEL_FILL
        r[n + 1].fill = ARTIST_FILL
    for i in range(1, n + 3):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions[get_column_letter(cols["track"] + 1)].width = 26
    ws.column_dimensions[get_column_letter(n + 1)].width = 20
    ws.column_dimensions[get_column_letter(n + 2)].width = 24
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_by_quarters(artist, uploaded_file, publish=True):
    """Создаёт или обновляет квартальные отчёты артиста из общего файла.

    Возвращает список (отчёт, создан_ли_новый) и число строк без периода."""
    header, cols, groups, bad = split_by_quarter(uploaded_file)
    results = []
    with transaction.atomic():
        for (year, quarter), rows in sorted(groups.items()):
            report = Report.objects.filter(artist=artist, year=year, quarter=quarter).first()
            created = report is None
            if created:
                report = Report(artist=artist, year=year, quarter=quarter, is_published=publish)
            report.royalty_rate = artist.royalty_rate
            data = build_detail_xlsx(artist.name, year, quarter, header, rows, cols, artist.royalty_rate)
            if report.detail_file:
                report.detail_file.delete(save=False)
            report.detail_file.save(f"detail_{year}_Q{quarter}.xlsx", ContentFile(data), save=False)
            report._parsed_rows = _collect(rows, cols)
            report.save()
            results.append((report, created))
    return results, bad
