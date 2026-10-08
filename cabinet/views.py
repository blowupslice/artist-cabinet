import mimetypes
import os
import re
from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

from django.conf import settings
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import update_session_auth_hash
from django.contrib.auth import views as auth_views
from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.db.models import Sum
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_GET

from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import BankDetailsForm, EmailLoginForm, PasswordChangeFormRu, SignedActForm
from .models import Artist, Contract, PlaylistPlacement, Release, Report, RoyaltyLine
from .stats import breakdown, pct_change, rate_per_1000, split


class LoginView(auth_views.LoginView):
    template_name = "cabinet/login.html"
    authentication_form = EmailLoginForm
    redirect_authenticated_user = True
    LOCK_SECONDS = 15 * 60

    def _key(self):
        email = (self.request.POST.get("username") or "").strip().lower()
        return f"login-fail:{email}"

    def post(self, request, *args, **kwargs):
        if cache.get(self._key(), 0) >= settings.LOGIN_MAX_FAILURES:
            form = self.get_form()
            form.errors.clear()
            form.add_error(None, "Слишком много неудачных попыток. Попробуйте через 15 минут.")
            return self.render_to_response(self.get_context_data(form=form))
        return super().post(request, *args, **kwargs)

    def form_invalid(self, form):
        key = self._key()
        cache.set(key, cache.get(key, 0) + 1, self.LOCK_SECONDS)
        return super().form_invalid(form)

    def form_valid(self, form):
        cache.delete(self._key())
        return super().form_valid(form)


# ---------------------------------------------------------------------------
# Какого артиста показываем
# ---------------------------------------------------------------------------

def get_artist(request):
    """Артист текущего пользователя. Администратор может открыть кабинет
    любого артиста (?as=<id>), чтобы увидеть его глазами артиста."""
    user = request.user
    if user.is_staff:
        as_id = request.GET.get("as")
        if as_id:
            request.session["view_as"] = as_id
        as_id = request.session.get("view_as")
        if as_id:
            artist = Artist.objects.filter(pk=as_id).select_related("user").first()
            if artist:
                return artist
    return Artist.objects.filter(user=user).select_related("user").first()


def artist_required(view):
    @login_required
    def wrapper(request, *args, **kwargs):
        artist = get_artist(request)
        if artist is None:
            if request.user.is_staff:
                return redirect("admin:cabinet_artist_changelist")
            return render(request, "cabinet/no_profile.html", status=403)
        request.artist = artist
        request.viewing_as = request.user.is_staff and artist.user_id != request.user.id
        return view(request, *args, **kwargs)
    wrapper.__name__ = view.__name__
    return wrapper


@login_required
def stop_view_as(request):
    request.session.pop("view_as", None)
    return redirect("admin:cabinet_artist_changelist")


def published_reports(artist):
    return artist.reports.filter(is_published=True).order_by("-year", "-quarter")


# ---------------------------------------------------------------------------
# Страницы
# ---------------------------------------------------------------------------

@artist_required
def dashboard(request):
    artist = request.artist
    reports = list(published_reports(artist))
    history = list(reversed(reports[:8]))
    max_total = max((r.total_amount for r in history), default=0) or 1
    for r in history:
        r.bar = float(r.total_amount / max_total * 100)
    to_sign = [r for r in reports if r.act_status == Report.ActStatus.TO_SIGN]
    total_all = sum((r.total_amount for r in reports), Decimal(0))
    from .acts import payout_overview
    payout = payout_overview(artist)
    releases = list(artist.releases.filter(is_published=True)[:20])
    upcoming = sorted([r for r in releases if r.is_upcoming], key=lambda r: r.release_date)
    latest_release = next((r for r in releases if not r.is_upcoming and r.status == Release.Status.RELEASED), None)
    playlists = list(artist.playlist_placements.all()[:5])
    fresh = None  # «Свежие прослушивания» пока скрыты
    todo = []
    if payout.report and payout.report.act_status == Report.ActStatus.TO_SIGN:
        todo.append(("act", payout.report))
    if not artist.has_bank_details:
        todo.append(("bank", None))
    return render(request, "cabinet/dashboard.html", {
        "artist": artist,
        "payout": payout,
        "todo": todo,
        "upcoming": upcoming[:2],
        "latest_release": latest_release,
        "playlists": playlists,
        "fresh": fresh,
        "contracts": artist.contracts.all(),
        "latest": reports[0] if reports else None,
        "history": history,
        "to_sign": to_sign,
        "total_all": total_all,
        "reports_count": len(reports),
        "currency": reports[0].currency if reports else settings.DEFAULT_CURRENCY,
        "nav": "dashboard",
    })


@artist_required
def finance(request, slug=None):
    artist = request.artist
    reports = list(published_reports(artist))
    if not reports:
        return render(request, "cabinet/finance.html", {"artist": artist, "reports": [], "nav": "finance"})

    report = reports[0]
    if slug:
        report = next((r for r in reports if r.slug == slug), None)
        if report is None:
            raise Http404("Отчёт не найден")

    lines = report.lines.all()
    by_platform = list(
        lines.values("platform").annotate(amount=Sum("amount"), quantity=Sum("quantity")).order_by("-amount")
    )
    by_track = list(
        lines.values("track").annotate(amount=Sum("amount"), quantity=Sum("quantity")).order_by("-amount")
    )
    pairs = {}
    for row in lines.values("track", "platform").annotate(amount=Sum("amount"), quantity=Sum("quantity")).order_by("-amount"):
        pairs.setdefault(row["track"], []).append(row)

    total = report.total_amount or Decimal(0)
    top = by_platform[0]["amount"] if by_platform else 1
    for p in by_platform:
        p["share"] = float(p["amount"] / total * 100) if total else 0
        p["bar"] = float(p["amount"] / top * 100) if top else 0
        p["rate"] = rate_per_1000(p["amount"], p["quantity"])
    for t in by_track:
        t["share"] = float(t["amount"] / total * 100) if total else 0
        t["platforms"] = pairs.get(t["track"], [])

    idx = reports.index(report)
    prev_report = reports[idx + 1] if idx + 1 < len(reports) else None
    delta = None
    if prev_report and prev_report.total_amount:
        delta = float((report.total_amount - prev_report.total_amount) / prev_report.total_amount * 100)

    countries = breakdown(lines, "country")
    leaders = {
        "track": by_track[0] if by_track else None,
        "platform": by_platform[0] if by_platform else None,
        "country": countries[0] if countries and any(c["key"] for c in countries) else None,
    }
    from .acts import compute
    act = compute(report)

    return render(request, "cabinet/finance.html", {
        "artist": artist,
        "reports": reports,
        "report": report,
        "leaders": leaders,
        "rights": split(lines, "rights_type"),
        "usage": split(lines, "usage_type"),
        "act": act,
        "act_form": SignedActForm(),
        "prev_report": prev_report,
        "delta": delta,
        "by_platform": by_platform,
        "by_track": by_track,
        "nav": "finance",
    })


def _prev_label(sel, prev_reports):
    if not prev_reports:
        return ""
    if not sel.isdigit():
        return "к " + prev_reports[0].period_short
    qs = sorted(r.quarter for r in prev_reports)
    year = prev_reports[0].year
    if len(qs) == 4:
        return f"к {year} году"
    if len(qs) == 1:
        return f"к Q{qs[0]} {year}"
    return f"к Q{qs[0]}–Q{qs[-1]} {year}"


@artist_required
def stats(request):
    artist = request.artist
    reports = list(published_reports(artist))
    from .streams import recent_summary
    # Блок «Свежие прослушивания» пока скрыт (данные в админке сохраняются)
    extra = {"fresh": None, "playlists": list(artist.playlist_placements.all()[:50])}
    if not reports:
        return render(request, "cabinet/stats.html", {"artist": artist, "reports": [], "nav": "stats", **extra})

    years = sorted({r.year for r in reports}, reverse=True)
    # «За всё время» пока убрано — по умолчанию открывается последний год
    options = [(str(y), f"{y} год") for y in years] + [(r.slug, r.period_short) for r in reports]
    sel = request.GET.get("p") or str(years[0])
    if sel not in {k for k, _ in options}:
        sel = str(years[0])

    prev_reports = None
    if sel == "all":
        selected = reports
        title = "За всё время"
    elif sel.isdigit():
        y = int(sel)
        selected = [r for r in reports if r.year == y]
        title = f"{y} год"
        # сравниваем с теми же кварталами прошлого года, а не с целым годом
        qs_sel = {r.quarter for r in selected}
        prev = [r for r in reports if r.year == y - 1 and r.quarter in qs_sel]
        prev_reports = prev or None
    else:
        idx = next(i for i, r in enumerate(reports) if r.slug == sel)
        selected = [reports[idx]]
        title = reports[idx].period_label
        prev_reports = [reports[idx + 1]] if idx + 1 < len(reports) else None

    lines = RoyaltyLine.objects.filter(report__in=selected)
    prev_lines = RoyaltyLine.objects.filter(report__in=prev_reports) if prev_reports else None

    total = sum((r.total_amount for r in selected), Decimal(0))
    streams = sum((r.total_quantity for r in selected), 0)
    prev_total = sum((r.total_amount for r in prev_reports), Decimal(0)) if prev_reports else None
    prev_streams = sum((r.total_quantity for r in prev_reports), 0) if prev_reports else None

    platforms = breakdown(lines, "platform", prev_lines)
    tracks = breakdown(lines, "track", prev_lines)
    countries = breakdown(lines, "country", prev_lines)
    has_countries = any(c["key"] for c in countries)

    # Динамика по кварталам (по возрастанию времени)
    timeline = list(reversed(reports))
    max_amount = max((r.total_amount for r in timeline), default=0) or 1
    max_streams = max((r.total_quantity for r in timeline), default=0) or 1
    selected_ids = {r.pk for r in selected}
    dynamics = []
    prev_r = None
    for r in timeline:
        dynamics.append({
            "report": r,
            "amount_bar": float(r.total_amount / max_amount * 100),
            "streams_bar": r.total_quantity / max_streams * 100,
            "rate": rate_per_1000(r.total_amount, r.total_quantity),
            "delta": pct_change(r.total_amount, prev_r.total_amount) if prev_r else None,
            "selected": sel != "all" and r.pk in selected_ids,
        })
        prev_r = r

    best_quarter = max(reports, key=lambda r: r.total_amount)
    return render(request, "cabinet/stats.html", {
        "artist": artist,
        "reports": reports,
        "options": options,
        "sel": sel,
        "title": title,
        "currency": reports[0].currency,
        "total": total,
        "streams": streams,
        "rate": rate_per_1000(total, streams),
        "delta_amount": pct_change(total, prev_total) if prev_total else None,
        "delta_streams": pct_change(streams, prev_streams) if prev_streams else None,
        "has_prev": bool(prev_reports),
        "prev_label": _prev_label(sel, prev_reports),
        "platforms": platforms,
        "tracks": tracks,
        "countries": countries if has_countries else [],
        "dynamics": dynamics,
        "best_quarter": best_quarter,
        "top_track": tracks[0] if tracks else None,
        "top_platform": platforms[0] if platforms else None,
        "nav": "stats",
        **extra,
    })


@artist_required
def account(request):
    artist = request.artist
    which = request.POST.get("form") if request.method == "POST" else None
    form = PasswordChangeFormRu(request.user, request.POST if which == "password" else None)
    bank_form = BankDetailsForm(request.POST if which == "bank" else None, instance=artist)
    if which and request.viewing_as:
        messages.error(request, "В режиме просмотра за артиста сохранять нельзя")
        return redirect("account")
    if which == "password" and form.is_valid():
        form.save()
        update_session_auth_hash(request, request.user)
        messages.success(request, "Пароль изменён")
        return redirect("account")
    if which == "bank" and bank_form.is_valid():
        a = bank_form.save(commit=False)
        a.bank_updated_at = timezone.now()
        a.save()
        messages.success(request, "Реквизиты сохранены. Выплаты придут на этот счёт.")
        nxt = request.GET.get("next") or ""
        return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else "account")
    return render(request, "cabinet/account.html", {
        "artist": artist, "form": form, "bank_form": bank_form, "nav": "account",
        "open_bank": which == "bank" or request.GET.get("tab") == "bank" or not artist.has_bank_details,
    })


@artist_required
@require_POST
def upload_signed_act(request, pk):
    report = get_object_or_404(Report, pk=pk, artist=request.artist, is_published=True)
    back = redirect("finance_period", report.slug)
    if request.viewing_as:
        messages.error(request, "В режиме просмотра за артиста загружать файлы нельзя")
        return back
    if report.act_status == Report.ActStatus.SIGNED:
        messages.error(request, "Этот акт уже принят — загружать заново не нужно")
        return back
    form = SignedActForm(request.POST, request.FILES)
    if not form.is_valid():
        messages.error(request, " ".join(form.errors.get("file", ["Не удалось загрузить файл"])))
        return back
    if report.signed_act_file:
        report.signed_act_file.delete(save=False)
    report.signed_act_file = form.cleaned_data["file"]
    report.act_status = Report.ActStatus.REVIEW
    report.signed_uploaded_at = timezone.now()
    report._from_artist = True
    report.save()
    messages.success(request, "Акт получили, спасибо! Менеджер проверит его и отправит выплату.")
    return back


@artist_required
def releases(request):
    artist = request.artist
    items = list(artist.releases.filter(is_published=True))
    upcoming = sorted([r for r in items if r.is_upcoming], key=lambda r: r.release_date)
    past = [r for r in items if not r.is_upcoming]
    return render(request, "cabinet/releases.html", {
        "artist": artist, "upcoming": upcoming, "past": past, "nav": "releases",
    })


# ---------------------------------------------------------------------------
# Скачивание файлов (только своих)
# ---------------------------------------------------------------------------

def _safe_name(text):
    text = re.sub(r"[^\w\-. ]+", "", text, flags=re.UNICODE).strip().replace(" ", "_")
    text = re.sub(r"_+", "_", text).strip("_")
    return text or "file"


def _send(field, download_name):
    if not field:
        raise Http404("Файл не загружен")
    try:
        fh = field.open("rb")
    except FileNotFoundError as exc:
        raise Http404("Файл не найден") from exc
    ext = os.path.splitext(field.name)[1].lower()
    ctype = mimetypes.guess_type(field.name)[0] or "application/octet-stream"
    return FileResponse(fh, as_attachment=True, filename=f"{download_name}{ext}", content_type=ctype)


@artist_required
@require_GET
def download_contract(request, pk):
    contract = get_object_or_404(Contract, pk=pk, artist=request.artist)
    return _send(contract.file, _safe_name(f"{contract}"))


@artist_required
@require_GET
def download_report_file(request, pk, kind):
    report = get_object_or_404(Report, pk=pk, artist=request.artist, is_published=True)
    base = _safe_name(f"{request.artist.name}_{report.period_short}")
    if kind == "detail":
        return _send(report.detail_file, f"Отчет_{base}")
    if kind == "act":
        if report.act_file:
            return _send(report.act_file, f"Акт_{base}")
        return _act_pdf_response(report, f"Акт_{base}")
    if kind == "signed-act":
        return _send(report.signed_act_file, f"Акт_подписанный_{base}")
    raise Http404


@artist_required
@require_GET
def export_summary(request, pk):
    """Сводка по кварталу в Excel: по трекам, по площадкам и трек × площадка."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    report = get_object_or_404(Report, pk=pk, artist=request.artist, is_published=True)
    lines = report.lines.all()
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F2937")

    def sheet(ws, title, headers, rows, widths):
        ws.title = title
        ws.append([f"{request.artist.name} — {report.period_label}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(headers)
        for c in ws[3]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=4):
            row[-1].number_format = "#,##0.00"
        ws.freeze_panes = "A4"

    cur = report.currency
    sheet(wb.active, "По трекам", ["Трек", "Прослушивания", f"Сумма, {cur}"],
          [(r["track"], r["q"], float(r["a"])) for r in
           lines.values("track").annotate(q=Sum("quantity"), a=Sum("amount")).order_by("-a")], [50, 16, 16])
    sheet(wb.create_sheet(), "По площадкам", ["Площадка", "Прослушивания", f"Сумма, {cur}"],
          [(r["platform"], r["q"], float(r["a"])) for r in
           lines.values("platform").annotate(q=Sum("quantity"), a=Sum("amount")).order_by("-a")], [30, 16, 16])
    sheet(wb.create_sheet(), "Трек × площадка", ["Трек", "Площадка", "Прослушивания", f"Сумма, {cur}"],
          [(r["track"], r["platform"], r["q"], float(r["a"])) for r in
           lines.values("track", "platform").annotate(q=Sum("quantity"), a=Sum("amount")).order_by("track", "-a")],
          [50, 30, 16, 16])

    buf = BytesIO()
    wb.save(buf)
    name = f"Сводка_{_safe_name(request.artist.name)}_{report.period_short.replace(' ', '_')}.xlsx"
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(name)}"
    return resp


def _act_pdf_response(report, name):
    from .acts import render_pdf
    pdf = render_pdf(report)
    resp = HttpResponse(pdf, content_type="application/pdf")
    resp["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(name + '.pdf')}"
    return resp


@staff_member_required
def admin_act_pdf(request, pk):
    report = get_object_or_404(Report, pk=pk)
    return _act_pdf_response(report, _safe_name(f"Акт_{report.artist.name}_{report.period_short}"))


@artist_required
@require_GET
def release_cover(request, pk):
    release = get_object_or_404(Release, pk=pk, artist=request.artist, is_published=True)
    if not release.cover:
        raise Http404
    try:
        fh = release.cover.open("rb")
    except FileNotFoundError as exc:
        raise Http404 from exc
    ctype = mimetypes.guess_type(release.cover.name)[0] or "image/jpeg"
    resp = FileResponse(fh, content_type=ctype)
    resp["Cache-Control"] = "private, max-age=86400"
    return resp


@staff_member_required
def protected_file(request, path):
    """Ссылки на файлы в админке. Доступны только администраторам."""
    from django.core.files.storage import default_storage
    if ".." in path or not default_storage.exists(path):
        raise Http404
    return FileResponse(default_storage.open(path, "rb"), as_attachment=False,
                        filename=os.path.basename(path))


def healthz(request):
    return HttpResponse("ok")
